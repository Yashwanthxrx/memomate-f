import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_report():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles
    COLOR_PRIMARY = RGBColor(217, 72, 59)      # Coral / Persimmon Red
    COLOR_DARK = RGBColor(26, 32, 44)          # Deep Slate
    COLOR_BODY = RGBColor(45, 55, 72)          # Charcoal Body
    COLOR_MUTED = RGBColor(113, 128, 150)      # Muted Gray

    # --- COVER / HEADER TITLE ---
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_sub = title_p.add_run("MONGODB STUDENT HACKATHON 2026 • TRACK 1: AGENTIC MEMORY\n")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(10)
    run_sub.font.bold = True
    run_sub.font.color.rgb = COLOR_PRIMARY

    run_title = title_p.add_run("Samantha (MemoMate)\nOS1 Intimate Diary & Memory Companion")
    run_title.font.name = "Georgia"
    run_title.font.size = Pt(26)
    run_title.font.bold = True
    run_title.font.color.rgb = COLOR_PRIMARY

    tagline_p = doc.add_paragraph()
    tagline_p.paragraph_format.space_after = Pt(20)
    run_tag = tagline_p.add_run("A Persistent, Privacy-First Agentic Memory Architecture Powered by MongoDB Atlas Vector Search, LangGraph StateGraph, and Google Gemini.")
    run_tag.font.name = "Arial"
    run_tag.font.size = Pt(12)
    run_tag.font.italic = True
    run_tag.font.color.rgb = COLOR_MUTED

    # Meta Info Card
    meta_table = doc.add_table(rows=1, cols=3)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    col_widths = [Inches(2.1), Inches(2.1), Inches(2.1)]
    for i, w in enumerate(col_widths):
        meta_table.columns[i].width = w

    meta_items = [
        ("PROJECT TRACK", "Agentic Persistent Memory"),
        ("DATABASE", "MongoDB Atlas Vector Search"),
        ("ORCHESTRATION", "LangGraph Multi-Step Graph")
    ]
    for i, (label, val) in enumerate(meta_items):
        cell = meta_table.rows[0].cells[i]
        set_cell_background(cell, "FBF4F1")
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r_lbl = p.add_run(f"{label}\n")
        r_lbl.font.size = Pt(8.5)
        r_lbl.font.bold = True
        r_lbl.font.color.rgb = COLOR_PRIMARY
        r_val = p.add_run(val)
        r_val.font.size = Pt(10)
        r_val.font.bold = True
        r_val.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(15)

    # --- SECTION 1: EXECUTIVE SUMMARY ---
    h1 = doc.add_heading("1. Executive Summary", level=1)
    h1.paragraph_format.space_before = Pt(14)
    h1.paragraph_format.space_after = Pt(6)
    for r in h1.runs:
        r.font.name = "Georgia"
        r.font.color.rgb = COLOR_PRIMARY

    p_exec = doc.add_paragraph(
        "Samantha (MemoMate) is an agentic diary and life-planning assistant inspired by Spike Jonze’s Oscar-winning film 'Her'. "
        "Standard conversational AI models operate with zero cross-session memory, forcing users to repeatedly restate habits, schedules, and preferences. "
        "Conversely, naive persistent bots accumulate stale, duplicate, or contradicting data without resolution, posing significant reliability and privacy hazards.\n\n"
        "Samantha overcomes these challenges by implementing an autonomous, lifecycle-aware memory engine. Leveraging a compiled LangGraph StateGraph, "
        "MongoDB Atlas Vector Search, and a fine-grained TTL index, Samantha extracts structured memory representations (Preferences, Facts, Goals, and Moods), "
        "audits conflict evolution, enforces automated redaction of sensitive credentials, and organically alters future conversations, recommendations, and schedules."
    )
    p_exec.style.font.name = "Arial"
    p_exec.style.font.size = Pt(10.5)

    # --- SECTION 2: HACKATHON CHALLENGE BRIEF ALIGNMENT ---
    h2 = doc.add_heading("2. Hackathon Brief Compliance Matrix", level=1)
    for r in h2.runs:
        r.font.name = "Georgia"
        r.font.color.rgb = COLOR_PRIMARY

    matrix_p = doc.add_paragraph(
        "Samantha was engineered directly to satisfy and exceed all required criteria established in the hackathon challenge prompt:"
    )

    table = doc.add_table(rows=7, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    widths = [Inches(1.8), Inches(2.7), Inches(2.0)]
    for i, w in enumerate(widths):
        table.columns[i].width = w

    headers = ["Challenge Requirement", "Samantha Architecture & Implementation", "Demonstration Scenario"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_background(cell, "D9483B")
        set_cell_margins(cell, top=140, bottom=140, left=140, right=140)
        p = cell.paragraphs[0]
        run = p.add_run(h)
        run.font.name = "Arial"
        run.font.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(255, 255, 255)

    data = [
        ("Useful Memory Creation", "Autonomous extraction into structured JSON schema (type, content, expires_on) with automatic 384-dimensional dense vector embeddings.", "Scenario 1: Form Memory\nDiary entry saves late-night study habit & tea preference."),
        ("Behavior-Changing Retrieval", "MongoDB Vector Search retrieves top-5 active memories. Semantic context is injected into generation, directly customizing future schedules.", "Scenario 2: Recall & Plan\nGenerates custom study routine anchored around known night habits."),
        ("Memory Inspection & Scores", "Transparent Memory Cortex displays category badges, expiry dates, and cosine relevance percentages influencing each reply.", "Expandable 'Memories Used' badges on every assistant message."),
        ("Conflict & Evolution (Hard Case)", "Vector similarity lookup + LLM conflict evaluation. Superseded memories transition to status='replaced' with replaced_by link.", "Scenario 3: Conflict / Replace\nMorning routine supersedes night habit with audit history."),
        ("Privacy Guard (Hard Case)", "Deterministic regex safety node blocks passwords, emails, phone numbers, and 6+ digit numbers before persistence.", "Scenario 4: Privacy Guard\nRedacts secret credentials to [SENSITIVE DATA REDACTED]."),
        ("Targeted Forget (Hard Case)", "Natural language forget command ('forget that I like...') performs cosine similarity search and permanently purges memory.", "Scenario 5: Vector Forget\nSemantic hard deletion from MongoDB Atlas collection.")
    ]

    for row_idx, row_data in enumerate(data, start=1):
        bg = "FBF4F1" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, text in enumerate(row_data):
            cell = table.rows[row_idx].cells[col_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=120, bottom=120, left=130, right=130)
            p = cell.paragraphs[0]
            run = p.add_run(text)
            run.font.name = "Arial"
            run.font.size = Pt(9)
            run.font.color.rgb = COLOR_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # --- SECTION 3: SYSTEM ARCHITECTURE & AGENT PIPELINE ---
    h3 = doc.add_heading("3. Agentic Pipeline Architecture", level=1)
    for r in h3.runs:
        r.font.name = "Georgia"
        r.font.color.rgb = COLOR_PRIMARY

    doc.add_paragraph(
        "Samantha is structured as a compiled LangGraph StateGraph engine, replacing chaotic monolithic prompts with deterministic, auditable multi-step agent nodes:"
    )

    nodes = [
        ("1. Privacy Guard Node", "First barrier of defense. Analyzes input text for email addresses, telephone numbers, and 6+ digit credentials. Flags and redacts sensitive patterns to prevent credential contamination."),
        ("2. Intent Router", "Conditional edge routing user inputs. If a forget pattern is identified ('forget...', 'delete memory...'), execution routes directly to the Forget Node; otherwise, it branches into the memory pipeline."),
        ("3. Memory Extractor Node", "Uses Google Gemini (with failover models) to extract atomic memories into four categories: Preference (habits/likes), Fact (life events), Goal (aspirations), and Mood (emotional states)."),
        ("4. Conflict Resolver Node (FR3 - Hard Case)", "Performs a vector search for the top-3 most similar active memories. An LLM evaluates whether the new memory 'replaces', 'duplicates', or is 'unrelated' to existing knowledge. Replaced memories are preserved with an audit pointer."),
        ("5. Semantic Recall Node (FR2)", "Executes MongoDB Atlas Vector Search ($vectorSearch stage) using cosine similarity, filtered by user_id and active status. Pulls the top-5 relevant memories and attaches similarity scores."),
        ("6. Response Generator Node", "Synthesizes recalled memories with conversation history to generate an empathetic, context-aware reply in Samantha's distinctive warm persona, subsequently persisting both the dialogue and memory provenance."),
        ("7. Forget Node (FR5)", "Executes targeted cosine similarity search against the user's active memories to pinpoint and delete the exact target memory without requiring exact string matching.")
    ]

    for title, desc in nodes:
        p_node = doc.add_paragraph()
        p_node.paragraph_format.space_before = Pt(3)
        p_node.paragraph_format.space_after = Pt(4)
        r_t = p_node.add_run(f"• {title}: ")
        r_t.font.name = "Arial"
        r_t.font.bold = True
        r_t.font.size = Pt(9.5)
        r_t.font.color.rgb = COLOR_PRIMARY
        r_d = p_node.add_run(desc)
        r_d.font.name = "Arial"
        r_d.font.size = Pt(9.5)
        r_d.font.color.rgb = COLOR_BODY

    # --- SECTION 4: MONGODB ATLAS INTEGRATION & DATA MODEL ---
    h4 = doc.add_heading("4. MongoDB Atlas Integration & Schema", level=1)
    for r in h4.runs:
        r.font.name = "Georgia"
        r.font.color.rgb = COLOR_PRIMARY

    doc.add_paragraph(
        "Samantha utilizes MongoDB Atlas as an enterprise document and vector store, taking advantage of hybrid query capabilities:"
    )

    doc.add_paragraph(
        "• Vector Search Index (memory_vector_index): Configured over the 384-dimensional dense embedding field using cosine similarity, "
        "incorporating pre-filter indexes on user_id, status, and memory type for high-efficiency scoped recall.\n"
        "• Time-To-Live (TTL) Index: Automatically purges transient memories once their expiration date passes. "
        "Moods expire after 72 hours (3 days), dated facts expire the following day, while preferences and goals remain permanent.\n"
        "• Document Schemas: Model both atomic memory documents and dialogue interaction documents with provenance tracking."
    )

    schema_text = (
        "Memory Document Schema:\n"
        "{\n"
        "  \"_id\": ObjectId(\"...\"),\n"
        "  \"user_id\": \"Theodore\",\n"
        "  \"type\": \"preference | fact | goal | mood\",\n"
        "  \"content\": \"Studies best late at night with green tea\",\n"
        "  \"embedding\": [0.0551, 0.0461, ... 384-dim float array],\n"
        "  \"source\": \"original user message\",\n"
        "  \"status\": \"active | replaced\",\n"
        "  \"replaced_by\": ObjectId or null,\n"
        "  \"created_at\": ISODate(\"...\"),\n"
        "  \"updated_at\": ISODate(\"...\"),\n"
        "  \"expires_at\": ISODate(\"...\") or null\n"
        "}"
    )
    p_code = doc.add_paragraph()
    p_code.paragraph_format.space_before = Pt(6)
    p_code.paragraph_format.space_after = Pt(12)
    r_code = p_code.add_run(schema_text)
    r_code.font.name = "Consolas"
    r_code.font.size = Pt(8.5)
    r_code.font.color.rgb = COLOR_DARK

    # --- SECTION 5: USER EXPERIENCE & CINEMATIC UI ---
    h5 = doc.add_heading("5. User Interface & The 'Her' Persona", level=1)
    for r in h5.runs:
        r.font.name = "Georgia"
        r.font.color.rgb = COLOR_PRIMARY

    doc.add_paragraph(
        "The interface was built to evoke the warm, retro-modern, intimate ambiance of the movie 'Her':\n"
        "• Breathing Consciousness Orb: A pulsing coral red dot in the header provides a visual representation of Samantha's presence.\n"
        "• Dual Atmosphere Modes: Seamless toggle between 'Midnight Coral' (deep slate obsidian #0D1117) and 'Clean Daylight' (studio paper #F5F7FA) with zero text-contrast degradation.\n"
        "• Ambient Soundtrack: Built-in tape-piano soundtrack player in the sidebar inspired by 'Song on the Beach'.\n"
        "• 1-Click Interactive Demo Flow: A multi-column test bench allowing judges to trigger all five hackathon challenge scenarios with a single click."
    )

    # --- SECTION 6: VERIFICATION & TEST SUITE ---
    h6 = doc.add_heading("6. Verification & Automated Test Results", level=1)
    for r in h6.runs:
        r.font.name = "Georgia"
        r.font.color.rgb = COLOR_PRIMARY

    doc.add_paragraph(
        "The complete LangGraph pipeline was verified using automated end-to-end integration tests (test_agent.py). "
        "All test suites executed with 100% pass rates:"
    )

    test_logs = (
        "[OK] LangGraph StateGraph compiled and instantiated\n"
        "[OK] test_privacy_guard_node: Redacted phone & email successfully\n"
        "[OK] test_intent_routing: Correctly routed forget vs memory pipeline\n"
        "[OK] Turn 1 completed: Memory saved -> ['Studies best late at night']\n"
        "[OK] Turn 2 completed: Recalled memories -> ['Studies best late at night']\n"
        "[OK] Turn 3 completed: Forget node successfully removed target memory\n"
        ">>> ALL LANGGRAPH AGENT TESTS PASSED! <<<"
    )
    p_tlog = doc.add_paragraph()
    r_tlog = p_tlog.add_run(test_logs)
    r_tlog.font.name = "Consolas"
    r_tlog.font.size = Pt(8.5)
    r_tlog.font.color.rgb = RGBColor(39, 110, 50)

    # --- SECTION 7: REPOSITORY & SUBMISSION DETAILS ---
    h7 = doc.add_heading("7. Submission & Repository Access", level=1)
    for r in h7.runs:
        r.font.name = "Georgia"
        r.font.color.rgb = COLOR_PRIMARY

    doc.add_paragraph(
        "• GitHub Repository: https://github.com/Yashwanthxrx/memomate-f\n"
        "• Submission ZIP: Samantha_MemoMate_Hackathon_Submission.zip (sanitized, portable, ~75 KB)\n"
        "• Active Branch: main\n"
        "• License: MIT Open Source License"
    )

    output_path = os.path.abspath("Samantha_MemoMate_Hackathon_Project_Report.docx")
    doc.save(output_path)
    print(f"Report successfully saved to: {output_path}")

if __name__ == "__main__":
    create_report()
